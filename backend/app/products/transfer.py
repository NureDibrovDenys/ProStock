"""Імпорт та експорт товарів у CSV / XLSX."""
from __future__ import annotations

import csv
import io
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.enums import UnitOfMeasure
from app.products import crud
from app.products.models import Category, Product
from app.products.schemas import ImportReport, ImportRowError, ProductImportRow

# Порядок колонок у шаблоні / експорті
EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("sku_code", "Артикул"),
    ("name", "Назва"),
    ("category", "Категорія"),
    ("unit_of_measure", "Одиниця виміру"),
    ("width", "Ширина, м"),
    ("length", "Довжина, м"),
    ("height", "Висота, м"),
    ("weight", "Вага, кг"),
    ("min_stock_level", "Мін. залишок"),
]

HEADER_ALIASES: dict[str, set[str]] = {
    "sku_code": {"sku_code", "sku", "code", "артикул", "код", "код_товару"},
    "name": {"name", "product_name", "назва", "найменування", "назва_товару"},
    "category": {"category", "category_name", "категорія", "категорія_товару"},
    "category_id": {"category_id", "категорія_id"},
    "unit_of_measure": {
        "unit_of_measure", "unit", "одиниця", "одиниця_виміру", "од_виміру", "од.виміру",
    },
    "width": {"width", "ширина"},
    "length": {"length", "довжина"},
    "height": {"height", "висота"},
    "weight": {"weight", "вага"},
    "min_stock_level": {
        "min_stock_level", "min_stock", "мін_залишок", "мінімальний_залишок",
        "мін._залишок",
    },
}

REQUIRED_COLUMNS = {"sku_code", "name", "unit_of_measure", "width", "length", "height", "weight"}

MAX_FILE_BYTES = settings.IMPORT_MAX_FILE_MB * 1024 * 1024


class ImportFileError(Exception):
    """Файл не читається або має некоректну структуру."""


# ── Читання файлу ──────────────────────────────────────────────────────────

def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\ufeff", "").strip().lower()
    for old, new in ((" ", "_"), ("-", "_"), ("/", "_")):
        text = text.replace(old, new)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


def _canonical_headers(raw_headers: list[object]) -> tuple[dict[int, str], set[str]]:
    mapping: dict[int, str] = {}
    seen: set[str] = set()
    for index, raw in enumerate(raw_headers):
        normalized = _normalize_header(raw)
        for canonical, aliases in HEADER_ALIASES.items():
            if normalized in aliases and canonical not in seen:
                mapping[index] = canonical
                seen.add(canonical)
                break
    return mapping, seen


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("Не вдалося визначити кодування файлу. Збережіть його як UTF-8 CSV.")


def _read_csv(content: bytes) -> list[list[object]]:
    text = _decode_csv(content)
    if not text.strip():
        raise ImportFileError("Файл порожній")

    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
    except csv.Error:
        pass

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [row for row in reader]


def _read_xlsx(content: bytes) -> list[list[object]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl кидає різні типи помилок
        raise ImportFileError("Не вдалося прочитати Excel-файл") from exc

    sheet = workbook.active
    if sheet is None:
        raise ImportFileError("У файлі немає жодного аркуша")

    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def read_table(filename: str, content: bytes) -> list[list[object]]:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(
            f"Файл завеликий: максимум {settings.IMPORT_MAX_FILE_MB} МБ"
        )

    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        return _read_csv(content)
    if suffix in {"xlsx", "xlsm"}:
        return _read_xlsx(content)
    raise ImportFileError("Підтримуються лише файли .csv та .xlsx")


# ── Валідація ──────────────────────────────────────────────────────────────

def _humanize_error(error: dict) -> str:
    """Перекладає типові помилки pydantic на зрозумілі українські тексти."""
    message = error.get("msg", "Некоректне значення")
    error_type = error.get("type", "")

    translations = {
        "missing": "Обовʼязкове поле",
        "string_too_short": "Значення занадто коротке",
        "string_too_long": "Значення занадто довге",
        "greater_than": "Значення має бути більше 0",
        "greater_than_equal": "Значення не може бути відʼємним",
        "decimal_parsing": "Некоректне число",
        "decimal_max_digits": "Занадто багато цифр у числі",
        "decimal_max_places": "Занадто багато знаків після коми",
        "int_parsing": "Некоректне ціле число",
    }
    if error_type in translations:
        return translations[error_type]
    # Власні ValueError з валідаторів уже українською
    return message.replace("Value error, ", "")


def _parse_category_id(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        parsed = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    return parsed if parsed > 0 else None


def _cell(row: list[object], index: int) -> object | None:
    return row[index] if index < len(row) else None


async def import_products(
    session: AsyncSession,
    filename: str,
    content: bytes,
) -> ImportReport:
    """Імпортує товари з файлу, повертаючи детальний звіт про помилки."""
    if not content:
        raise ImportFileError("Файл порожній")

    table = read_table(filename, content)
    # Відкидаємо повністю порожні рядки
    table = [row for row in table if any(str(cell).strip() for cell in row if cell is not None)]
    if not table:
        raise ImportFileError("Файл не містить даних")

    mapping, seen_headers = _canonical_headers(table[0])
    missing = REQUIRED_COLUMNS - seen_headers
    if "category" not in seen_headers and "category_id" not in seen_headers:
        missing.add("category")
    if missing:
        names = {
            "sku_code": "Артикул (sku_code)",
            "name": "Назва (name)",
            "category": "Категорія (category)",
            "unit_of_measure": "Одиниця виміру (unit_of_measure)",
            "width": "Ширина (width)",
            "length": "Довжина (length)",
            "height": "Висота (height)",
            "weight": "Вага (weight)",
        }
        pretty = ", ".join(names.get(column, column) for column in sorted(missing))
        raise ImportFileError(f"У файлі бракує обовʼязкових колонок: {pretty}")

    data_rows = table[1:]
    if len(data_rows) > settings.IMPORT_MAX_ROWS:
        raise ImportFileError(
            f"Забагато рядків: максимум {settings.IMPORT_MAX_ROWS}"
        )

    errors: list[ImportRowError] = []
    valid: list[tuple[int, ProductImportRow, str | None, int | None]] = []
    seen_skus: dict[str, int] = {}

    for offset, row in enumerate(data_rows, start=2):
        record: dict[str, object] = {}
        for index, canonical in mapping.items():
            record[canonical] = _cell(row, index)

        raw_sku = record.get("sku_code")
        sku = str(raw_sku).strip() if raw_sku is not None else None
        category_name = record.get("category")
        category_name = str(category_name).strip() if category_name is not None else None

        if sku:
            if sku in seen_skus:
                errors.append(
                    ImportRowError(
                        row=offset,
                        sku_code=sku,
                        field="sku_code",
                        message=f"Дублікат артикулу в файлі (рядок {seen_skus[sku]})",
                    )
                )
                continue
            seen_skus[sku] = offset

        payload = {
            "sku_code": raw_sku,
            "name": record.get("name"),
            "unit_of_measure": record.get("unit_of_measure"),
            "width": record.get("width"),
            "length": record.get("length"),
            "height": record.get("height"),
            "weight": record.get("weight"),
            "min_stock_level": record.get("min_stock_level"),
        }

        try:
            parsed = ProductImportRow.model_validate(payload)
        except ValidationError as exc:
            for error in exc.errors():
                field = str(error["loc"][0]) if error.get("loc") else None
                errors.append(
                    ImportRowError(
                        row=offset,
                        sku_code=sku,
                        field=field,
                        message=_humanize_error(error),
                    )
                )
            continue

        if not category_name and _parse_category_id(record.get("category_id")) is None:
            errors.append(
                ImportRowError(
                    row=offset,
                    sku_code=parsed.sku_code,
                    field="category",
                    message="Вкажіть категорію або її ідентифікатор",
                )
            )
            continue

        valid.append(
            (
                offset,
                parsed,
                category_name or None,
                _parse_category_id(record.get("category_id")),
            )
        )

    created = 0
    updated = 0
    categories_created = 0
    category_cache: dict[str, Category] = {}

    for offset, parsed, category_name, category_id in valid:
        category: Category | None = None
        try:
            if category_name:
                key = category_name.lower()
                category = category_cache.get(key)
                if category is None:
                    category = await crud.find_category_by_name(session, category_name)
                    if category is None:
                        category = Category(name=category_name[:150])
                        session.add(category)
                        await session.flush()
                        categories_created += 1
                    category_cache[key] = category
            elif category_id is not None:
                category = await session.get(Category, category_id)

            if category is None:
                errors.append(
                    ImportRowError(
                        row=offset,
                        sku_code=parsed.sku_code,
                        field="category",
                        message="Категорію не знайдено",
                    )
                )
                continue

            async with session.begin_nested():
                existing = await crud.get_product_by_sku(session, parsed.sku_code)
                if existing is None:
                    session.add(
                        Product(
                            sku_code=parsed.sku_code,
                            name=parsed.name,
                            category_id=category.id,
                            unit_of_measure=parsed.unit_of_measure,
                            width=parsed.width,
                            length=parsed.length,
                            height=parsed.height,
                            weight=parsed.weight,
                            min_stock_level=parsed.min_stock_level,
                        )
                    )
                    created += 1
                else:
                    existing.name = parsed.name
                    existing.category_id = category.id
                    existing.unit_of_measure = parsed.unit_of_measure
                    existing.width = parsed.width
                    existing.length = parsed.length
                    existing.height = parsed.height
                    existing.weight = parsed.weight
                    existing.min_stock_level = parsed.min_stock_level
                    updated += 1
        except SQLAlchemyError:
            errors.append(
                ImportRowError(
                    row=offset,
                    sku_code=parsed.sku_code,
                    field=None,
                    message="Не вдалося зберегти рядок у базі даних",
                )
            )

    if created or updated or categories_created:
        await session.commit()
    else:
        await session.rollback()

    return ImportReport(
        filename=filename,
        total=len(data_rows),
        created=created,
        updated=updated,
        failed=len(errors),
        categories_created=categories_created,
        errors=errors,
    )


# ── Експорт ────────────────────────────────────────────────────────────────

def _row_values(product: Product) -> list[object]:
    return [
        product.sku_code,
        product.name,
        product.category.name if product.category else "",
        product.unit_of_measure.value if isinstance(product.unit_of_measure, UnitOfMeasure) else product.unit_of_measure,
        _decimal_str(product.width),
        _decimal_str(product.length),
        _decimal_str(product.height),
        _decimal_str(product.weight),
        _decimal_str(product.min_stock_level),
    ]


def _decimal_str(value: Decimal | None) -> str:
    if value is None:
        return ""
    return format(Decimal(value).normalize(), "f")


def build_csv(products: list[Product]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([key for key, _ in EXPORT_COLUMNS])
    for product in products:
        writer.writerow(_row_values(product))
    # BOM, щоб Excel правильно відкрив кирилицю
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def build_xlsx(products: list[Product]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Товари"

    header_fill = PatternFill("solid", fgColor="656B46")
    header_font = Font(bold=True, color="FFFFFF")

    for column_index, (key, title) in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=key)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(column_index)].width = max(14, len(title) + 4)

    for row_index, product in enumerate(products, start=2):
        for column_index, value in enumerate(_row_values(product), start=1):
            sheet.cell(row=row_index, column=column_index, value=value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(1, len(products) + 1)}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_template(fmt: str) -> bytes:
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
        writer.writerow([key for key, _ in EXPORT_COLUMNS])
        return ("\ufeff" + buffer.getvalue()).encode("utf-8")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Товари"
    header_fill = PatternFill("solid", fgColor="656B46")
    header_font = Font(bold=True, color="FFFFFF")
    for column_index, (key, title) in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=key)
        cell.fill = header_fill
        cell.font = header_font
        sheet.column_dimensions[get_column_letter(column_index)].width = max(16, len(title) + 4)

    instructions = workbook.create_sheet("Інструкція")
    instructions.column_dimensions["A"].width = 100
    instructions["A1"] = "Заповніть аркуш «Товари». Обовʼязкові колонки: sku_code, name, unit_of_measure, width, length, height, weight."
    instructions["A2"] = "Одиниці виміру: pcs, kg, pound, box, pallet, other (або шт, кг, коробка, палета)."
    instructions["A3"] = "Категорія: назва існуючої категорії або нова — вона створиться автоматично."
    instructions["A4"] = "Існуючі артикули (sku_code) оновлюються, нові — створюються."

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(extension: str) -> str:
    return f"products_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{extension}"
